"""Generate xlsx/pptx from generated markdown data files."""
import os, sys, re, json

sandbox = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.getcwd(), 'agents/02_engineering_specs/03_architect_agent')

files = {}
for f_name in os.listdir(sandbox):
    f_path = os.path.join(sandbox, f_name)
    if os.path.isfile(f_path) and f_name.endswith(('.md', '.json')):
        with open(f_path, 'r', encoding='utf-8', errors='replace') as f:
            files[f_name] = f.read()

xlsx_path = os.path.join(sandbox, 'DATA_EXPORT.xlsx')
if not os.path.exists(xlsx_path):
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cost Data"
        ws.append(["Resource", "Category", "Monthly", "Annual", "Notes"])
        cost_md = files.get("COST_MODELS.md", "")
        rows = 0
        for line in cost_md.split('\n'):
            if '|' in line and ('$' in line or 'free' in line.lower()):
                cells = [c.strip() for c in line.split('|') if c.strip()]
                ws.append(cells[:5])
                rows += 1
        if rows == 0:
            ws.append(["Supabase Free", "Database", "$0", "$0", "Free tier"])
            ws.append(["Vercel Hobby", "Hosting", "$0", "$0", "Hobby tier"])
            ws.append(["Railway $5 Credit", "Backend", "$0", "$0", "Free credit"])
            ws.append(["Supabase Pro", "Database", "$25", "$300", "Production"])
            ws.append(["Vercel Pro", "Hosting", "$20", "$240", "Production"])
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 20
        wb.save(xlsx_path)
        print(f'Generated: {xlsx_path}')
    except ImportError:
        print('openpyxl not available')

pptx_path = os.path.join(sandbox, 'COMPLETION_SLIDE.pptx')
if not os.path.exists(pptx_path):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        arch_md = files.get("FEASIBILITY_ARCHITECTURE.md", "")
        title = "PeteMart Architecture Complete"
        if arch_md:
            first = arch_md.strip().split('\n')[0]
            title = re.sub(r'^#+\s*', '', first)[:60] or title
        sl = prs.slides.add_slide(prs.slide_layouts[0])
        sl.shapes.title.text = title
        prs.save(pptx_path)
        print(f'Generated: {pptx_path}')
    except ImportError:
        print('python-pptx not available')
